from pptx import Presentation
prs = Presentation(r"C:\Users\USER\Downloads\intern\w3\sprint2_submission\PPT Slides\Sprint2_Financial_Ratio_Engine.pptx")
for i, slide in enumerate(prs.slides):
    print(f"--- Slide {i+1} ---")
    for shape in slide.shapes:
        if hasattr(shape, "text"):
            print(shape.text)
