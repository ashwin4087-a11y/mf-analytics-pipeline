from pptx import Presentation
from pptx.util import Inches

prs = Presentation()
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Sprint 2 — Financial Ratio Engine"
subtitle.text = "Substantially Complete / Blocked by Source Data Limitation"

bullet_slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(bullet_slide_layout)
shapes = slide.shapes
title_shape = shapes.title
body_shape = shapes.placeholders[1]
title_shape.text = "SPRINT 2 STATUS"
tf = body_shape.text_frame
tf.text = "SUBSTANTIALLY COMPLETE\nBLOCKED BY SOURCE DATA LIMITATION"

p = tf.add_paragraph()
p.text = "51 computed KPIs"
p.level = 1
p = tf.add_paragraph()
p.text = "66 total financial_ratios columns"
p.level = 1
p = tf.add_paragraph()
p.text = "31/31 tests passing"
p.level = 1
p = tf.add_paragraph()
p.text = "1,070 legitimate integrated rows"
p.level = 1
p = tf.add_paragraph()
p.text = "1,100+ required threshold"
p.level = 1
p = tf.add_paragraph()
p.text = "91 source-dependent rows not integrable"
p.level = 1
p = tf.add_paragraph()
p.text = "8 missing master company IDs"
p.level = 1
p = tf.add_paragraph()
p.text = "Data Governance Note: Missing master records were not fabricated or inferred."
p.level = 0

prs.save(r"C:\Users\USER\Downloads\intern\w3\sprint2_submission\PPT Slides\Sprint2_Financial_Ratio_Engine.pptx")
